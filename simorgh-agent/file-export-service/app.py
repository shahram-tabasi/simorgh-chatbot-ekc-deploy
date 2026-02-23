"""
File Export Service
=====================
Generate Excel, Word, and PDF files from structured data.
Used by the agent to produce deliverables from project analysis.

Endpoints:
  POST /export/excel    - Generate Excel file
  POST /export/word     - Generate Word document
  POST /export/pdf      - Generate PDF report
  GET  /download/{id}   - Download generated file
  GET  /health          - Health check
"""

import io
import logging
import os
import uuid
from datetime import datetime
from pathlib import Path
from typing import Optional, List, Dict, Any

from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = FastAPI(title="File Export Service", version="1.0.0")

EXPORT_DIR = Path(os.getenv("EXPORT_DIR", "/app/exports"))
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

_exports: Dict[str, Dict] = {}


class TableData(BaseModel):
    headers: List[str]
    rows: List[List[str]]
    sheet_name: str = "Sheet1"


class ExcelExportRequest(BaseModel):
    project_id: str
    filename: str = Field("export.xlsx")
    title: str = Field("Export")
    tables: List[TableData]


class WordExportRequest(BaseModel):
    project_id: str
    filename: str = Field("report.docx")
    title: str = Field("Project Report")
    sections: List[Dict[str, Any]] = Field(
        ..., description="List of {heading, content, level} dicts"
    )
    tables: List[TableData] = []


class PdfExportRequest(BaseModel):
    project_id: str
    filename: str = Field("report.pdf")
    title: str = Field("Project Report")
    content: str = Field(..., description="Markdown or plain text content")


class ExportResponse(BaseModel):
    export_id: str
    filename: str
    format: str
    size: int
    download_url: str


@app.get("/health")
async def health():
    return {"status": "healthy", "service": "file-export"}


@app.post("/export/excel", response_model=ExportResponse)
async def export_excel(req: ExcelExportRequest):
    """Generate an Excel file from tabular data."""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        for table in req.tables:
            ws = wb.create_sheet(title=table.sheet_name[:31])

            # Header row
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=11)

            for col, header in enumerate(table.headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            # Data rows
            for row_idx, row_data in enumerate(table.rows, 2):
                for col_idx, value in enumerate(row_data, 1):
                    ws.cell(row=row_idx, column=col_idx, value=value)

            # Auto-width
            for col in ws.columns:
                max_len = max((len(str(cell.value or "")) for cell in col), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

        export_id = str(uuid.uuid4())
        filepath = EXPORT_DIR / f"{export_id}_{req.filename}"
        wb.save(str(filepath))

        size = filepath.stat().st_size
        _exports[export_id] = {
            "path": str(filepath),
            "filename": req.filename,
            "format": "xlsx",
            "size": size,
            "created_at": datetime.utcnow().isoformat(),
        }

        return ExportResponse(
            export_id=export_id,
            filename=req.filename,
            format="xlsx",
            size=size,
            download_url=f"/download/{export_id}",
        )

    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")
    except Exception as e:
        logger.error(f"Excel export failed: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/export/word", response_model=ExportResponse)
async def export_word(req: WordExportRequest):
    """Generate a Word document from sections and tables."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Title
        title_para = doc.add_heading(req.title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M')}")
        doc.add_paragraph(f"Project: {req.project_id}")
        doc.add_paragraph("")

        # Sections
        for section in req.sections:
            level = section.get("level", 1)
            heading = section.get("heading", "Section")
            content = section.get("content", "")

            doc.add_heading(heading, level=min(level, 4))
            for para_text in content.split("\n"):
                if para_text.strip():
                    doc.add_paragraph(para_text.strip())

        # Tables
        for table_data in req.tables:
            doc.add_heading(table_data.sheet_name, level=2)
            table = doc.add_table(
                rows=1 + len(table_data.rows),
                cols=len(table_data.headers),
                style="Table Grid",
            )

            # Header
            for i, header in enumerate(table_data.headers):
                cell = table.rows[0].cells[i]
                cell.text = header
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True

            # Rows
            for row_idx, row_data in enumerate(table_data.rows):
                for col_idx, value in enumerate(row_data):
                    if col_idx < len(table_data.headers):
                        table.rows[row_idx + 1].cells[col_idx].text = str(value)

        export_id = str(uuid.uuid4())
        filepath = EXPORT_DIR / f"{export_id}_{req.filename}"
        doc.save(str(filepath))

        size = filepath.stat().st_size
        _exports[export_id] = {
            "path": str(filepath),
            "filename": req.filename,
            "format": "docx",
            "size": size,
            "created_at": datetime.utcnow().isoformat(),
        }

        return ExportResponse(
            export_id=export_id,
            filename=req.filename,
            format="docx",
            size=size,
            download_url=f"/download/{export_id}",
        )

    except ImportError:
        raise HTTPException(status_code=500, detail="python-docx not installed")
    except Exception as e:
        logger.error(f"Word export failed: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/export/pdf", response_model=ExportResponse)
async def export_pdf(req: PdfExportRequest):
    """Generate a PDF from text/markdown content."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.units import inch

        export_id = str(uuid.uuid4())
        filepath = EXPORT_DIR / f"{export_id}_{req.filename}"

        doc = SimpleDocTemplate(str(filepath), pagesize=A4)
        styles = getSampleStyleSheet()

        # Custom styles
        title_style = ParagraphStyle(
            "CustomTitle", parent=styles["Title"], fontSize=18, spaceAfter=20
        )
        body_style = ParagraphStyle(
            "CustomBody", parent=styles["Normal"], fontSize=10, leading=14
        )

        story = []
        story.append(Paragraph(req.title, title_style))
        story.append(Paragraph(
            f"Project: {req.project_id} | Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M')}",
            styles["Italic"],
        ))
        story.append(Spacer(1, 0.3 * inch))

        # Convert content to paragraphs
        for line in req.content.split("\n"):
            line = line.strip()
            if not line:
                story.append(Spacer(1, 0.1 * inch))
            elif line.startswith("# "):
                story.append(Paragraph(line[2:], styles["Heading1"]))
            elif line.startswith("## "):
                story.append(Paragraph(line[3:], styles["Heading2"]))
            elif line.startswith("### "):
                story.append(Paragraph(line[4:], styles["Heading3"]))
            elif line.startswith("- "):
                story.append(Paragraph(f"&bull; {line[2:]}", body_style))
            else:
                # Escape XML special characters for reportlab
                safe_line = (
                    line.replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;")
                )
                story.append(Paragraph(safe_line, body_style))

        doc.build(story)

        size = filepath.stat().st_size
        _exports[export_id] = {
            "path": str(filepath),
            "filename": req.filename,
            "format": "pdf",
            "size": size,
            "created_at": datetime.utcnow().isoformat(),
        }

        return ExportResponse(
            export_id=export_id,
            filename=req.filename,
            format="pdf",
            size=size,
            download_url=f"/download/{export_id}",
        )

    except ImportError:
        raise HTTPException(status_code=500, detail="reportlab not installed")
    except Exception as e:
        logger.error(f"PDF export failed: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/download/{export_id}")
async def download_file(export_id: str):
    """Download a generated export file."""
    if export_id not in _exports:
        raise HTTPException(status_code=404, detail="Export not found")

    info = _exports[export_id]
    filepath = Path(info["path"])
    if not filepath.exists():
        raise HTTPException(status_code=404, detail="File no longer available")

    return FileResponse(
        path=str(filepath),
        filename=info["filename"],
        media_type="application/octet-stream",
    )


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8025)
